"""
DOCX 转 MD 转换器
使用 python-docx 提取文本和结构
"""

from docx import Document
from pathlib import Path
from typing import List

from .base import BaseConverter


class DOCXToMDConverter(BaseConverter):
    """DOCX 转 Markdown 转换器"""

    input_extensions = [".docx"]
    output_extension = ".md"

    def convert(self, input_path: str, output_path: str) -> bool:
        """
        将 DOCX 转换为 Markdown

        Args:
            input_path: DOCX 文件路径
            output_path: MD 文件路径

        Returns:
            bool: 转换是否成功
        """
        self.validate(input_path)

        doc = Document(input_path)
        lines = []

        for para in doc.paragraphs:
            style_name = para.style.name.lower() if para.style else ""

            # 标题处理
            if "heading" in style_name:
                level = para.style.name
                # 提取标题级别数字
                for i in range(1, 7):
                    if str(i) in para.style.name:
                        level = i
                        break
                else:
                    level = 1

                text = para.text.strip()
                if text:
                    lines.append(f"{'#' * level} {text}\n")
                continue

            # 列表处理
            if para.style and ("list" in style_name or "bullet" in style_name):
                text = para.text.strip()
                if text:
                    lines.append(f"- {text}")
                continue

            # 普通段落
            text = para.text.strip()
            if text:
                lines.append(f"{text}\n")

        # 处理表格
        for table in doc.tables:
            lines.append("\n| ")
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                lines.append(" | ".join(cells) + " |\n")
            lines.append("\n")

        # 写入文件
        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text("\n".join(lines), encoding="utf-8")

        return True
